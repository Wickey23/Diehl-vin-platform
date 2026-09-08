from __future__ import annotations

import json
import os
import subprocess
import sys
import threading
import webbrowser
from pathlib import Path

from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse
from pydantic import BaseModel
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
CONFIG = ROOT / 'local_dtna_config.json'
DTNA_RUNTIME = ROOT / 'dtna_runtime.py'
PORT = 8770

app = FastAPI(title='Diehl Local DTNA Writer', version='1.0')


def load_config() -> dict:
    if not CONFIG.exists():
        return {'workbook': '', 'sheet': ''}
    try:
        data = json.loads(CONFIG.read_text(encoding='utf-8'))
    except Exception:
        return {'workbook': '', 'sheet': ''}
    return {
        'workbook': str(data.get('workbook') or ''),
        'sheet': str(data.get('sheet') or ''),
    }


def save_config(workbook: Path, sheet: str) -> None:
    CONFIG.write_text(json.dumps({'workbook': str(workbook), 'sheet': sheet}, indent=2), encoding='utf-8')


def workbook_sheets(path: Path) -> list[str]:
    if not path.exists():
        raise RuntimeError('Workbook does not exist.')
    keep_vba = path.suffix.lower() == '.xlsm'
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=keep_vba)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def choose_workbook_native() -> str:
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    try:
        selected = filedialog.askopenfilename(
            title='Choose the Excel workbook DTNA should write to',
            filetypes=[
                ('Excel workbooks', '*.xlsx *.xlsm'),
                ('Excel workbook', '*.xlsx'),
                ('Macro-enabled Excel workbook', '*.xlsm'),
            ],
        )
    finally:
        root.destroy()
    return selected


class SheetSelection(BaseModel):
    sheet: str


def validate_config() -> tuple[Path, str, list[str]]:
    cfg = load_config()
    raw = cfg.get('workbook', '').strip()
    sheet = cfg.get('sheet', '').strip()
    if not raw:
        raise HTTPException(400, 'Choose a workbook first.')
    path = Path(os.path.expandvars(raw)).expanduser().resolve()
    if not path.exists():
        raise HTTPException(400, f'Workbook not found: {path}')
    if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise HTTPException(400, 'Workbook must be .xlsx or .xlsm.')
    try:
        sheets = workbook_sheets(path)
    except Exception as exc:
        raise HTTPException(400, f'Could not read workbook sheets: {exc}') from exc
    if not sheet:
        raise HTTPException(400, 'Choose an existing sheet first.')
    if sheet not in sheets:
        raise HTTPException(400, f'Sheet "{sheet}" no longer exists in the workbook.')
    return path, sheet, sheets


def launch_dtna(login_only: bool = False) -> None:
    path, sheet, _ = validate_config()
    if not DTNA_RUNTIME.exists():
        raise HTTPException(500, 'dtna_runtime.py is missing.')

    env = os.environ.copy()
    env['DIEHL_DTNA_WORKBOOK'] = str(path)
    env['DIEHL_DTNA_SHEET'] = sheet

    args = [sys.executable, str(DTNA_RUNTIME)]
    if login_only:
        args.append('--login-only')

    flags = getattr(subprocess, 'CREATE_NEW_CONSOLE', 0)
    subprocess.Popen(args, cwd=str(ROOT), env=env, creationflags=flags)


@app.get('/', response_class=HTMLResponse)
def home():
    return HTMLResponse(HTML)


@app.get('/api/config')
def get_config():
    cfg = load_config()
    raw = cfg.get('workbook', '').strip()
    sheets: list[str] = []
    exists = False
    if raw:
        path = Path(os.path.expandvars(raw)).expanduser()
        exists = path.exists()
        if exists:
            try:
                sheets = workbook_sheets(path)
            except Exception:
                sheets = []
    return {
        'ok': True,
        'workbook': raw,
        'sheet': cfg.get('sheet', ''),
        'exists': exists,
        'sheets': sheets,
    }


@app.post('/api/select-workbook')
def select_workbook():
    selected = choose_workbook_native()
    if not selected:
        return {'ok': False, 'cancelled': True}
    path = Path(selected).resolve()
    if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise HTTPException(400, 'Choose an .xlsx or .xlsm workbook.')
    try:
        sheets = workbook_sheets(path)
    except Exception as exc:
        raise HTTPException(400, f'Could not open workbook: {exc}') from exc
    if not sheets:
        raise HTTPException(400, 'The selected workbook has no worksheets.')

    current = load_config()
    current_sheet = current.get('sheet', '')
    sheet = current_sheet if current_sheet in sheets else sheets[0]
    save_config(path, sheet)
    return {'ok': True, 'workbook': str(path), 'sheet': sheet, 'sheets': sheets}


@app.post('/api/select-sheet')
def select_sheet(body: SheetSelection):
    cfg = load_config()
    raw = cfg.get('workbook', '').strip()
    if not raw:
        raise HTTPException(400, 'Choose a workbook first.')
    path = Path(raw).expanduser().resolve()
    try:
        sheets = workbook_sheets(path)
    except Exception as exc:
        raise HTTPException(400, f'Could not read workbook: {exc}') from exc
    if body.sheet not in sheets:
        raise HTTPException(400, 'That sheet does not exist in the selected workbook.')
    save_config(path, body.sheet)
    return {'ok': True, 'workbook': str(path), 'sheet': body.sheet, 'sheets': sheets}


@app.post('/api/open-login')
def open_login():
    launch_dtna(login_only=True)
    return {'ok': True, 'message': 'DTNA login opened in a local browser window.'}


@app.post('/api/run-dtna')
def run_dtna():
    path, sheet, _ = validate_config()
    launch_dtna(login_only=False)
    return {
        'ok': True,
        'message': 'DTNA sync started in a separate console/window.',
        'workbook': str(path),
        'sheet': sheet,
    }


@app.get('/api/verify')
def verify():
    path, sheet, sheets = validate_config()
    return {
        'ok': True,
        'workbook': str(path),
        'sheet': sheet,
        'sheetExists': sheet in sheets,
    }


HTML = r'''<!doctype html>
<html>
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>Diehl Local DTNA Writer</title>
<style>
:root{font-family:Inter,Segoe UI,Arial,sans-serif;color:#182230;background:#f4f7fb}*{box-sizing:border-box}body{margin:0}.wrap{max-width:980px;margin:48px auto;padding:0 20px}.hero{background:#153b61;color:#fff;border-radius:18px;padding:30px 34px;margin-bottom:22px}.hero h1{font-size:34px;margin:0 0 8px}.hero p{margin:0;color:#dce8f5}.card{background:#fff;border:1px solid #dce3eb;border-radius:16px;padding:24px;margin-bottom:18px;box-shadow:0 4px 18px rgba(16,24,40,.05)}.label{font-size:12px;font-weight:800;letter-spacing:.08em;color:#456b91;text-transform:uppercase;margin-bottom:8px}.path{font-family:Consolas,monospace;background:#f8fafc;border:1px solid #e3e8ef;padding:12px;border-radius:9px;word-break:break-all;margin:10px 0 14px}select,button{font:inherit}select{width:100%;padding:11px;border:1px solid #cbd5e1;border-radius:9px;background:white;margin-top:8px}button{border:0;border-radius:9px;padding:11px 16px;font-weight:700;cursor:pointer}.primary{background:#1463d6;color:#fff}.secondary{background:#eaf1f8;color:#183b61}.danger{background:#b42318;color:#fff}.row{display:flex;gap:10px;flex-wrap:wrap}.status{margin-top:14px;padding:12px;border-radius:9px;background:#f5f7fa;color:#475467}.status.good{background:#ecfdf3;color:#067647}.status.bad{background:#fef3f2;color:#b42318}.big{padding:14px 20px;font-size:16px}.muted{color:#667085;font-size:14px}.steps{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}@media(max-width:720px){.steps{grid-template-columns:1fr}}.step{background:#f8fafc;border:1px solid #e5eaf0;border-radius:12px;padding:15px}.step b{display:block;margin-bottom:5px}
</style>
</head>
<body>
<div class="wrap">
  <div class="hero"><h1>Local DTNA Writer</h1><p>Choose the exact Excel workbook and existing worksheet that this PC should update.</p></div>

  <div class="card">
    <div class="label">1 · Workbook</div>
    <div id="path" class="path">No workbook selected</div>
    <button class="primary" onclick="chooseWorkbook()">Choose Workbook</button>
    <p class="muted">Supports existing .xlsx and .xlsm files. The program writes back to the same file you select.</p>
  </div>

  <div class="card">
    <div class="label">2 · Existing Sheet</div>
    <select id="sheet" onchange="saveSheet()"><option value="">Choose a workbook first</option></select>
    <p class="muted">Only sheets already inside the workbook are listed. The DTNA run will not create a new sheet in localhost mode.</p>
  </div>

  <div class="card">
    <div class="label">3 · DTNA</div>
    <div class="row">
      <button class="secondary big" onclick="openLogin()">Open / Refresh DTNA Login</button>
      <button class="primary big" onclick="runDtna()">Run DTNA + Write Excel</button>
    </div>
    <div id="status" class="status">Choose a workbook and sheet.</div>
  </div>

  <div class="steps">
    <div class="step"><b>Exact file target</b><span class="muted">No OneDrive discovery or master-workbook search is used by this localhost run.</span></div>
    <div class="step"><b>Existing sheet only</b><span class="muted">DTNA clears and replaces the data area in the worksheet you select.</span></div>
    <div class="step"><b>Saved choice</b><span class="muted">Your workbook and sheet selection is remembered on this PC for the next run.</span></div>
  </div>
</div>
<script>
const pathEl=document.getElementById('path'),sheetEl=document.getElementById('sheet'),statusEl=document.getElementById('status');
function status(msg,kind=''){statusEl.textContent=msg;statusEl.className='status '+kind}
async function json(url,opts={}){const r=await fetch(url,{...opts,headers:{'content-type':'application/json',...(opts.headers||{})}});const d=await r.json().catch(()=>({}));if(!r.ok)throw new Error(d.detail||d.error||'Request failed');return d}
function render(d){pathEl.textContent=d.workbook||'No workbook selected';sheetEl.innerHTML='';if(d.sheets?.length){for(const s of d.sheets){const o=document.createElement('option');o.value=s;o.textContent=s;o.selected=s===d.sheet;sheetEl.appendChild(o)}}else{const o=document.createElement('option');o.textContent='Choose a workbook first';o.value='';sheetEl.appendChild(o)}if(d.workbook&&d.sheet)status(`Ready: ${d.sheet} in ${d.workbook}`,'good')}
async function load(){try{render(await json('/api/config'))}catch(e){status(e.message,'bad')}}
async function chooseWorkbook(){status('Opening Windows file picker…');try{const d=await json('/api/select-workbook',{method:'POST'});if(d.cancelled){status('Workbook selection cancelled.');return}render(d)}catch(e){status(e.message,'bad')}}
async function saveSheet(){if(!sheetEl.value)return;try{const d=await json('/api/select-sheet',{method:'POST',body:JSON.stringify({sheet:sheetEl.value})});render(d)}catch(e){status(e.message,'bad')}}
async function openLogin(){status('Opening DTNA login…');try{const d=await json('/api/open-login',{method:'POST'});status(d.message,'good')}catch(e){status(e.message,'bad')}}
async function runDtna(){status('Validating target and starting DTNA…');try{const d=await json('/api/run-dtna',{method:'POST'});status(`Started. DTNA will write to: ${d.workbook} → ${d.sheet}`,'good')}catch(e){status(e.message,'bad')}}
load();
</script>
</body>
</html>'''


def open_browser() -> None:
    webbrowser.open_new_tab(f'http://127.0.0.1:{PORT}/')


if __name__ == '__main__':
    import uvicorn
    threading.Timer(1.2, open_browser).start()
    uvicorn.run(app, host='127.0.0.1', port=PORT, log_level='warning')
