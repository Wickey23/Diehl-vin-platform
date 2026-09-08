from __future__ import annotations

import getpass
import json
import os
import re
import socket
from pathlib import Path

SOURCE_FOLDER_NAME = 'Diehl VIN Platform Sources'


def _safe_piece(value: str) -> str:
    value = re.sub(r'[^A-Za-z0-9._-]+', '_', str(value or '').strip())
    return value.strip('._-') or 'UNKNOWN'


def source_workbook_name() -> str:
    computer = _safe_piece(os.environ.get('COMPUTERNAME') or socket.gethostname())
    user = _safe_piece(getpass.getuser())
    return f'DIEHL-VIN-SOURCE_{computer}_{user}.xlsx'


WORKBOOK_NAME = source_workbook_name()


def _unique_paths(values):
    seen = set()
    out = []
    for value in values:
        if not value:
            continue
        try:
            path = Path(os.path.expandvars(str(value))).expanduser()
            key = os.path.normcase(os.path.abspath(str(path)))
        except Exception:
            continue
        if key in seen:
            continue
        seen.add(key)
        out.append(path)
    return out


def onedrive_roots() -> list[Path]:
    home = Path.home()
    candidates = [
        os.environ.get('OneDriveCommercial'),
        os.environ.get('OneDrive'),
        os.environ.get('OneDriveConsumer'),
    ]
    try:
        candidates.extend(str(p) for p in home.glob('OneDrive*') if p.is_dir())
    except Exception:
        pass
    return [p for p in _unique_paths(candidates) if p.exists() and p.is_dir()]


def preferred_onedrive_root() -> Path:
    roots = onedrive_roots()
    if not roots:
        raise RuntimeError(
            'No synced OneDrive folder was detected on this PC. Start/sign in to OneDrive, '
            'then run START DIEHL VIN again.'
        )

    commercial = os.environ.get('OneDriveCommercial')
    if commercial:
        p = Path(os.path.expandvars(commercial)).expanduser()
        if p.exists() and p.is_dir():
            return p

    for root in roots:
        text = str(root).casefold()
        if 'diehl' in text or 'truck world' in text:
            return root
    return roots[0]


def source_folder() -> Path:
    folder = preferred_onedrive_root() / SOURCE_FOLDER_NAME
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def expected_source_workbook() -> Path:
    return source_folder() / WORKBOOK_NAME


def _create_source_workbook(path: Path) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    dtna = wb.active
    dtna.title = 'DTNA'
    dtna.append([
        'VIN', 'inServiceDate', 'serialNo', 'leadSerialNo', 'changeCount',
        'changeNotes', 'lastChangeTime'
    ])

    vin = wb.create_sheet('VIN In-Service')
    vin.append([
        'VIN', 'Verification Status', 'In-Service Status', 'In-Service Date',
        'Mileage', 'Customer Result', 'Customer Name', 'Registered Customer Name',
        'Registered Customer Account', 'Ordered Customer Name', 'Last Updated', 'Updated By'
    ])

    info = wb.create_sheet('Source Info')
    info.append(['Field', 'Value'])
    info.append(['Computer', os.environ.get('COMPUTERNAME') or socket.gethostname()])
    info.append(['Windows User', getpass.getuser()])
    info.append(['Source Workbook', path.name])
    info.append(['Source Folder', str(path.parent)])
    info.append(['Purpose', 'Per-computer Diehl VIN source workbook'])

    wb.save(path)


def find_shared_workbook(cached_path: str | Path | None = None) -> Path:
    """Return this computer's own OneDrive source workbook.

    The legacy function name is retained so existing DTNA/VIN code does not need
    to change. Each PC writes only to its own workbook. A separate master workbook
    may pull every DIEHL-VIN-SOURCE_*.xlsx file from SOURCE_FOLDER_NAME.
    """
    expected = expected_source_workbook()
    expected_norm = os.path.normcase(os.path.abspath(str(expected)))

    if cached_path:
        try:
            cached = Path(os.path.expandvars(str(cached_path))).expanduser()
            cached_norm = os.path.normcase(os.path.abspath(str(cached)))
            if cached_norm == expected_norm and cached.exists() and cached.is_file():
                return cached
        except Exception:
            pass

    if not expected.exists():
        _create_source_workbook(expected)
    return expected


def load_cached_path(config_path: Path) -> str:
    try:
        data = json.loads(config_path.read_text(encoding='utf-8')) if config_path.exists() else {}
        return str(data.get('masterWorkbook') or '').strip()
    except Exception:
        return ''
